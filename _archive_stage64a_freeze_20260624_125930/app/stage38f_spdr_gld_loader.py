#!/usr/bin/env python3
"""
Stage38F SPDR GLD daily holdings/flow loader.

Purpose
-------
Read the SPDR GLD Historical Archive XLSX discovered by Stage38F source audit,
normalize daily GLD holdings/flow data, write it into SQLite, build simple daily
ETF-flow features, and anti-lookahead join those features to XAUUSD H1 bars.

This is data-foundation only. It does not create trading signals, backtests,
orders, paper-live actions, or Stage39 artifacts.

Default input expected from prior source audit:
    data/etf/gold/source_audit/raw/spdr_historical_archive.xlsx

Main output tables:
    stage38f_spdr_gld_daily
    stage38f_spdr_gld_features
    stage38f_spdr_gld_h1_joined
    stage38f_spdr_gld_loader_audit
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sqlite3
from dataclasses import dataclass, asdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from statistics import mean, pstdev
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    from openpyxl import load_workbook
except Exception as exc:  # pragma: no cover
    raise SystemExit(
        "ERROR: openpyxl is required. Install with: python3 -m pip install openpyxl"
    ) from exc

DAILY_TABLE = "stage38f_spdr_gld_daily"
FEATURE_TABLE = "stage38f_spdr_gld_features"
JOINED_TABLE = "stage38f_spdr_gld_h1_joined"
AUDIT_TABLE = "stage38f_spdr_gld_loader_audit"

SPDR_DEFAULT_FILE = "spdr_historical_archive.xlsx"
SPDR_DEFAULT_SHEET = "US GLD Historical Archive"

TIMESTAMP_CANDIDATES = [
    "utc_time",
    "ts_utc",
    "timestamp_utc",
    "datetime_utc",
    "bar_ts_utc",
    "time_utc",
    "open_time_utc",
    "bar_time_utc",
    "open_time",
    "source_time",
    "timestamp",
    "datetime",
    "time",
    "ts",
]

PRICE_COLS = ["open", "high", "low", "close"]

HEADER_ALIASES = {
    "date": "observation_date",
    "closing price": "closing_price",
    "ounces of gold per share": "ounces_gold_per_share",
    "nav/share at 10:30am nyt": "nav_share_1030_nyt",
    "indicative price per share at 4:15pm nyt": "indicative_price_1615_nyt",
    "mid point of bid/ask spread at 4:15pm nyt": "bid_ask_mid_1615_nyt",
    "premium/discount of gld mid point vs indicative value of gld at 4:15pm nyt": "premium_discount_mid_vs_indicative",
    "daily share volume": "daily_share_volume",
    "total ounces of gold in the trust": "total_ounces_gold_trust",
    "tonnes of gold": "tonnes_gold",
    "total net asset value in the trust": "total_nav_trust",
}


@dataclass
class Audit:
    status: str
    decision: str
    workbook_path: str
    sheet_name: str
    raw_rows_seen: int
    parsed_rows: int
    holiday_or_non_numeric_rows_skipped: int
    duplicate_date_count: int
    daily_rows_written: int
    feature_rows_written: int
    h1_bars_total: int
    h1_joined_count: int
    h1_unjoined_count: int
    lookahead_violation_count: int
    warning_count: int
    note_count: int
    min_observation_date: Optional[str]
    max_observation_date: Optional[str]
    min_available_from_utc: Optional[str]
    max_available_from_utc: Optional[str]
    json_report: str
    md_report: str
    daily_table: str = DAILY_TABLE
    feature_table: str = FEATURE_TABLE
    joined_table: str = JOINED_TABLE
    audit_table: str = AUDIT_TABLE


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def norm_header(x: Any) -> str:
    if x is None:
        return ""
    s = str(x).strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def to_float(x: Any) -> Optional[float]:
    if x is None:
        return None
    if isinstance(x, (int, float)) and not isinstance(x, bool):
        if math.isfinite(float(x)):
            return float(x)
        return None
    s = str(x).strip()
    if not s:
        return None
    if s.lower() in {"us holiday", "holiday", "n/a", "na", "nan", "-", "--"}:
        return None
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg = True
        s = s[1:-1]
    s = s.replace(",", "").replace("$", "").replace("%", "").strip()
    try:
        val = float(s)
    except Exception:
        return None
    if not math.isfinite(val):
        return None
    return -val if neg else val


def parse_date_cell(x: Any) -> Optional[date]:
    if x is None:
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x).strip()
    if not s or s.lower() in {"date", "us holiday", "holiday"}:
        return None
    for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            pass
    return None


def iso_date(d: date) -> str:
    return d.isoformat()


def available_from(d: date) -> str:
    # Conservative anti-lookahead default: use previous daily observation only from next UTC day.
    return datetime(d.year, d.month, d.day, tzinfo=timezone.utc).replace(
        hour=0, minute=0, second=0, microsecond=0
    ).__add__(timedelta(days=1)).isoformat().replace("+00:00", "Z")


def parse_iso_ts(x: str) -> datetime:
    s = str(x).strip()
    if s.endswith("Z"):
        s = s[:-1] + "+00:00"
    dt = datetime.fromisoformat(s)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def find_workbook(raw_dir: Path, explicit: Optional[str]) -> Path:
    if explicit:
        p = Path(explicit)
        if not p.exists():
            raise FileNotFoundError(f"Workbook not found: {p}")
        return p
    p = raw_dir / SPDR_DEFAULT_FILE
    if p.exists():
        return p
    candidates = sorted(raw_dir.glob("*spdr*.xlsx")) + sorted(raw_dir.glob("*.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"No XLSX workbook found in {raw_dir}")
    return candidates[0]


def find_header_and_columns(ws) -> Tuple[int, Dict[str, int], List[str]]:
    best: Tuple[int, int, Dict[str, int], List[str]] = (-1, -1, {}, [])
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        headers = [norm_header(c) for c in row]
        colmap: Dict[str, int] = {}
        score = 0
        for j, h in enumerate(headers):
            if h in HEADER_ALIASES:
                colmap[HEADER_ALIASES[h]] = j
                score += 4
            elif "date" == h:
                colmap["observation_date"] = j
                score += 4
            elif "tonnes" in h and "gold" in h:
                colmap["tonnes_gold"] = j
                score += 4
            elif "total ounces" in h and "trust" in h:
                colmap["total_ounces_gold_trust"] = j
                score += 4
        if "observation_date" in colmap and "tonnes_gold" in colmap:
            score += 20
        if score > best[0]:
            best = (score, idx, colmap, [str(c) if c is not None else "" for c in row])
    score, header_row, colmap, preview = best
    if score < 20 or "observation_date" not in colmap or "tonnes_gold" not in colmap:
        raise RuntimeError(f"Cannot detect SPDR GLD header row. best_score={score}, preview={preview}")
    return header_row, colmap, preview


def parse_spdr_workbook(path: Path, sheet_name_arg: Optional[str]) -> Tuple[str, List[Dict[str, Any]], Dict[str, Any]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet_name_arg and sheet_name_arg in wb.sheetnames:
        ws = wb[sheet_name_arg]
    elif SPDR_DEFAULT_SHEET in wb.sheetnames:
        ws = wb[SPDR_DEFAULT_SHEET]
    else:
        ws = wb[wb.sheetnames[0]]

    header_row, colmap, header_preview = find_header_and_columns(ws)
    rows: List[Dict[str, Any]] = []
    seen: set[str] = set()
    dupes = 0
    skipped = 0
    raw_seen = 0

    for ridx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if ridx <= header_row:
            continue
        raw_seen += 1
        d = parse_date_cell(row[colmap["observation_date"]] if colmap["observation_date"] < len(row) else None)
        if d is None:
            skipped += 1
            continue
        tonnes = to_float(row[colmap["tonnes_gold"]] if colmap["tonnes_gold"] < len(row) else None)
        ounces = to_float(row[colmap.get("total_ounces_gold_trust", -1)] if colmap.get("total_ounces_gold_trust", 10**9) < len(row) else None)
        if tonnes is None and ounces is None:
            skipped += 1
            continue
        key = iso_date(d)
        if key in seen:
            dupes += 1
            continue
        seen.add(key)

        rec: Dict[str, Any] = {
            "observation_date": key,
            "available_from_utc": available_from(d),
            "closing_price": None,
            "ounces_gold_per_share": None,
            "nav_share_1030_nyt": None,
            "indicative_price_1615_nyt": None,
            "bid_ask_mid_1615_nyt": None,
            "premium_discount_mid_vs_indicative": None,
            "daily_share_volume": None,
            "total_ounces_gold_trust": ounces,
            "tonnes_gold": tonnes,
            "total_nav_trust": None,
            "source_file": path.name,
            "source_sheet": ws.title,
            "loaded_at_utc": utc_now_iso(),
        }
        for field, idx in colmap.items():
            if field == "observation_date":
                continue
            if idx < len(row):
                rec[field] = to_float(row[idx])
        # If tonnes missing but ounces present, approximate metric tonnes.
        if rec.get("tonnes_gold") is None and rec.get("total_ounces_gold_trust") is not None:
            rec["tonnes_gold"] = float(rec["total_ounces_gold_trust"]) / 32150.746568627
        if rec.get("total_ounces_gold_trust") is None and rec.get("tonnes_gold") is not None:
            rec["total_ounces_gold_trust"] = float(rec["tonnes_gold"]) * 32150.746568627
        rows.append(rec)

    rows.sort(key=lambda r: r["observation_date"])
    meta = {
        "sheet_name": ws.title,
        "header_row_1based": header_row,
        "header_preview": header_preview,
        "raw_rows_seen": raw_seen,
        "holiday_or_non_numeric_rows_skipped": skipped,
        "duplicate_date_count": dupes,
    }
    return ws.title, rows, meta


def rolling_z(values: Sequence[Optional[float]], idx: int, window: int) -> Optional[float]:
    v = values[idx]
    if v is None:
        return None
    start = max(0, idx - window + 1)
    sample = [x for x in values[start : idx + 1] if x is not None]
    if len(sample) < max(20, min(window, 60)):
        return None
    sd = pstdev(sample)
    if sd == 0:
        return 0.0
    return (v - mean(sample)) / sd


def chg(values: Sequence[Optional[float]], idx: int, lag: int) -> Optional[float]:
    if idx - lag < 0:
        return None
    a = values[idx]
    b = values[idx - lag]
    if a is None or b is None:
        return None
    return a - b


def pct_chg(values: Sequence[Optional[float]], idx: int, lag: int) -> Optional[float]:
    if idx - lag < 0:
        return None
    a = values[idx]
    b = values[idx - lag]
    if a is None or b is None or b == 0:
        return None
    return (a / b - 1.0) * 100.0


def flow_state(chg_5d: Optional[float], chg_20d: Optional[float]) -> str:
    # Conservative coarse buckets; tuned only for annotation, not strategy.
    if chg_5d is None and chg_20d is None:
        return "ETF_FLOW_UNKNOWN"
    c5 = chg_5d if chg_5d is not None else 0.0
    c20 = chg_20d if chg_20d is not None else 0.0
    if c5 >= 5.0 and c20 >= 10.0:
        return "ETF_STRONG_INFLOW"
    if c5 > 1.0 or c20 > 5.0:
        return "ETF_INFLOW"
    if c5 <= -5.0 and c20 <= -10.0:
        return "ETF_STRONG_OUTFLOW"
    if c5 < -1.0 or c20 < -5.0:
        return "ETF_OUTFLOW"
    return "ETF_FLOW_FLAT"


def build_features(daily_rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    tonnes = [r.get("tonnes_gold") for r in daily_rows]
    ounces = [r.get("total_ounces_gold_trust") for r in daily_rows]
    out: List[Dict[str, Any]] = []
    for i, r in enumerate(daily_rows):
        t1 = chg(tonnes, i, 1)
        t5 = chg(tonnes, i, 5)
        t20 = chg(tonnes, i, 20)
        row = {
            "observation_date": r["observation_date"],
            "available_from_utc": r["available_from_utc"],
            "tonnes_gold": r.get("tonnes_gold"),
            "total_ounces_gold_trust": r.get("total_ounces_gold_trust"),
            "closing_price": r.get("closing_price"),
            "total_nav_trust": r.get("total_nav_trust"),
            "gld_tonnes_chg_1d": t1,
            "gld_tonnes_chg_5d": t5,
            "gld_tonnes_chg_20d": t20,
            "gld_tonnes_pct_chg_5d": pct_chg(tonnes, i, 5),
            "gld_tonnes_pct_chg_20d": pct_chg(tonnes, i, 20),
            "gld_ounces_chg_5d": chg(ounces, i, 5),
            "gld_holding_zscore_156d": rolling_z(tonnes, i, 156),
            "gld_flow_state": flow_state(t5, t20),
            "loaded_at_utc": utc_now_iso(),
        }
        out.append(row)
    return out


def drop_and_create_tables(con: sqlite3.Connection) -> None:
    cur = con.cursor()
    for t in [DAILY_TABLE, FEATURE_TABLE, JOINED_TABLE, AUDIT_TABLE]:
        cur.execute(f'drop table if exists "{t}"')

    cur.execute(f'''
    create table "{DAILY_TABLE}" (
        observation_date text primary key,
        available_from_utc text not null,
        closing_price real,
        ounces_gold_per_share real,
        nav_share_1030_nyt real,
        indicative_price_1615_nyt real,
        bid_ask_mid_1615_nyt real,
        premium_discount_mid_vs_indicative real,
        daily_share_volume real,
        total_ounces_gold_trust real,
        tonnes_gold real,
        total_nav_trust real,
        source_file text,
        source_sheet text,
        loaded_at_utc text
    )
    ''')

    cur.execute(f'''
    create table "{FEATURE_TABLE}" (
        observation_date text primary key,
        available_from_utc text not null,
        tonnes_gold real,
        total_ounces_gold_trust real,
        closing_price real,
        total_nav_trust real,
        gld_tonnes_chg_1d real,
        gld_tonnes_chg_5d real,
        gld_tonnes_chg_20d real,
        gld_tonnes_pct_chg_5d real,
        gld_tonnes_pct_chg_20d real,
        gld_ounces_chg_5d real,
        gld_holding_zscore_156d real,
        gld_flow_state text,
        loaded_at_utc text
    )
    ''')

    cur.execute(f'''
    create table "{JOINED_TABLE}" (
        bar_ts_utc text primary key,
        source text,
        symbol text,
        timeframe text,
        open real,
        high real,
        low real,
        close real,
        spread real,
        gld_observation_date text,
        gld_available_from_utc text,
        gld_tonnes_gold real,
        gld_total_ounces_gold_trust real,
        gld_tonnes_chg_1d real,
        gld_tonnes_chg_5d real,
        gld_tonnes_chg_20d real,
        gld_tonnes_pct_chg_5d real,
        gld_tonnes_pct_chg_20d real,
        gld_holding_zscore_156d real,
        gld_flow_state text
    )
    ''')

    cur.execute(f'''
    create table "{AUDIT_TABLE}" (
        audit_id integer primary key autoincrement,
        created_at_utc text,
        status text,
        decision text,
        workbook_path text,
        sheet_name text,
        raw_rows_seen integer,
        parsed_rows integer,
        holiday_or_non_numeric_rows_skipped integer,
        duplicate_date_count integer,
        daily_rows_written integer,
        feature_rows_written integer,
        h1_bars_total integer,
        h1_joined_count integer,
        h1_unjoined_count integer,
        lookahead_violation_count integer,
        warning_count integer,
        note_count integer,
        min_observation_date text,
        max_observation_date text,
        min_available_from_utc text,
        max_available_from_utc text,
        json_report text,
        md_report text
    )
    ''')
    con.commit()


def insert_rows(con: sqlite3.Connection, table: str, rows: List[Dict[str, Any]]) -> None:
    if not rows:
        return
    keys = list(rows[0].keys())
    sql = f'insert into "{table}" ({", ".join([f"\"{k}\"" for k in keys])}) values ({", ".join(["?" for _ in keys])})'
    con.executemany(sql, [[r.get(k) for k in keys] for r in rows])
    con.commit()


def table_cols(con: sqlite3.Connection, table: str) -> List[str]:
    return [r[1] for r in con.execute(f'pragma table_info("{table}")').fetchall()]


def detect_col(cols: List[str], candidates: Sequence[str], role: str) -> str:
    lower = {c.lower(): c for c in cols}
    for c in candidates:
        if c.lower() in lower:
            return lower[c.lower()]
    raise RuntimeError(f"Cannot detect {role}. Tried={candidates}; existing={cols}")


def load_h1_bars(
    con: sqlite3.Connection,
    bars_table: str,
    source: str,
    symbol: str,
    timeframe_aliases: Sequence[str],
) -> List[Dict[str, Any]]:
    cols = table_cols(con, bars_table)
    ts_col = detect_col(cols, TIMESTAMP_CANDIDATES, "bar timestamp column")
    select_cols = [ts_col]
    for c in ["source", "symbol", "timeframe", "open", "high", "low", "close", "spread"]:
        if c in cols:
            select_cols.append(c)
    placeholders = ",".join(["?" for _ in timeframe_aliases])
    sql = f'''
    select {", ".join([f'"{c}"' for c in select_cols])}
    from "{bars_table}"
    where source = ? and symbol = ? and timeframe in ({placeholders})
    order by "{ts_col}"
    '''
    rows = con.execute(sql, [source, symbol, *timeframe_aliases]).fetchall()
    out: List[Dict[str, Any]] = []
    for tup in rows:
        d = dict(zip(select_cols, tup))
        ts = str(d[ts_col]).replace("+00:00", "Z")
        out.append({
            "bar_ts_utc": ts,
            "source": d.get("source", source),
            "symbol": d.get("symbol", symbol),
            "timeframe": d.get("timeframe", timeframe_aliases[0]),
            "open": to_float(d.get("open")),
            "high": to_float(d.get("high")),
            "low": to_float(d.get("low")),
            "close": to_float(d.get("close")),
            "spread": to_float(d.get("spread")),
        })
    return out


def join_features_to_bars(bars: List[Dict[str, Any]], features: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], int]:
    feat_sorted = sorted(features, key=lambda r: r["available_from_utc"])
    feat_times = [parse_iso_ts(r["available_from_utc"]) for r in feat_sorted]
    joined: List[Dict[str, Any]] = []
    fi = -1
    violations = 0
    for b in bars:
        bt = parse_iso_ts(b["bar_ts_utc"])
        while fi + 1 < len(feat_times) and feat_times[fi + 1] <= bt:
            fi += 1
        if fi < 0:
            continue
        f = feat_sorted[fi]
        if parse_iso_ts(f["available_from_utc"]) > bt:
            violations += 1
        joined.append({
            "bar_ts_utc": b["bar_ts_utc"],
            "source": b["source"],
            "symbol": b["symbol"],
            "timeframe": b["timeframe"],
            "open": b.get("open"),
            "high": b.get("high"),
            "low": b.get("low"),
            "close": b.get("close"),
            "spread": b.get("spread"),
            "gld_observation_date": f["observation_date"],
            "gld_available_from_utc": f["available_from_utc"],
            "gld_tonnes_gold": f.get("tonnes_gold"),
            "gld_total_ounces_gold_trust": f.get("total_ounces_gold_trust"),
            "gld_tonnes_chg_1d": f.get("gld_tonnes_chg_1d"),
            "gld_tonnes_chg_5d": f.get("gld_tonnes_chg_5d"),
            "gld_tonnes_chg_20d": f.get("gld_tonnes_chg_20d"),
            "gld_tonnes_pct_chg_5d": f.get("gld_tonnes_pct_chg_5d"),
            "gld_tonnes_pct_chg_20d": f.get("gld_tonnes_pct_chg_20d"),
            "gld_holding_zscore_156d": f.get("gld_holding_zscore_156d"),
            "gld_flow_state": f.get("gld_flow_state"),
        })
    return joined, violations


def write_reports(report_dir: Path, audit: Audit, daily_rows: List[Dict[str, Any]], feature_rows: List[Dict[str, Any]]) -> None:
    report_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "audit": asdict(audit),
        "latest_daily_rows": daily_rows[-5:],
        "latest_feature_rows": feature_rows[-5:],
    }
    Path(audit.json_report).write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")

    md = []
    md.append("# Stage38F SPDR GLD Loader Audit")
    md.append("")
    md.append(f"- status: `{audit.status}`")
    md.append(f"- decision: `{audit.decision}`")
    md.append(f"- workbook: `{audit.workbook_path}`")
    md.append(f"- sheet: `{audit.sheet_name}`")
    md.append(f"- parsed rows: `{audit.parsed_rows}`")
    md.append(f"- daily rows written: `{audit.daily_rows_written}`")
    md.append(f"- feature rows written: `{audit.feature_rows_written}`")
    md.append(f"- H1 joined: `{audit.h1_joined_count}/{audit.h1_bars_total}`")
    md.append(f"- lookahead violations: `{audit.lookahead_violation_count}`")
    md.append("")
    md.append("## Latest GLD feature rows")
    md.append("")
    for r in feature_rows[-8:]:
        md.append(
            f"- {r['observation_date']}: tonnes={r.get('tonnes_gold')}, "
            f"chg_5d={r.get('gld_tonnes_chg_5d')}, chg_20d={r.get('gld_tonnes_chg_20d')}, "
            f"state={r.get('gld_flow_state')}"
        )
    Path(audit.md_report).write_text("\n".join(md) + "\n", encoding="utf-8")


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", default="data/local/xauusd_local_store.sqlite")
    ap.add_argument("--bars-table", default="bars")
    ap.add_argument("--raw-dir", default="data/etf/gold/source_audit/raw")
    ap.add_argument("--workbook", default=None)
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--source", default="amarkets_mt5")
    ap.add_argument("--symbol", default="XAUUSD")
    ap.add_argument("--h1-timeframes", default="H1,1h")
    ap.add_argument("--reports-dir", default="data/reports/stage38f_spdr_gld_loader")
    args = ap.parse_args(argv)

    db = Path(args.db)
    raw_dir = Path(args.raw_dir)
    report_dir = Path(args.reports_dir)
    report_dir.mkdir(parents=True, exist_ok=True)
    json_report = str(report_dir / "stage38f_spdr_gld_loader.json")
    md_report = str(report_dir / "stage38f_spdr_gld_loader.md")

    workbook = find_workbook(raw_dir, args.workbook)
    sheet_name, daily_rows, meta = parse_spdr_workbook(workbook, args.sheet)
    feature_rows = build_features(daily_rows)

    con = sqlite3.connect(str(db))
    drop_and_create_tables(con)
    insert_rows(con, DAILY_TABLE, daily_rows)
    insert_rows(con, FEATURE_TABLE, feature_rows)

    aliases = [x.strip() for x in args.h1_timeframes.split(",") if x.strip()]
    h1_bars = load_h1_bars(con, args.bars_table, args.source, args.symbol, aliases)
    joined_rows, violations = join_features_to_bars(h1_bars, feature_rows)
    insert_rows(con, JOINED_TABLE, joined_rows)

    warnings: List[str] = []
    notes: List[str] = []
    if not daily_rows:
        warnings.append("no_spdr_daily_rows_parsed")
    if not feature_rows:
        warnings.append("no_spdr_feature_rows_built")
    if len(joined_rows) < len(h1_bars):
        notes.append(f"h1_unjoined_before_first_available_gld={len(h1_bars) - len(joined_rows)}")
    if meta.get("holiday_or_non_numeric_rows_skipped", 0):
        notes.append(f"holiday_or_non_numeric_rows_skipped={meta['holiday_or_non_numeric_rows_skipped']}")
    if meta.get("duplicate_date_count", 0):
        warnings.append(f"duplicate_date_count={meta['duplicate_date_count']}")
    if violations:
        warnings.append(f"lookahead_violation_count={violations}")

    if daily_rows and feature_rows and joined_rows and not violations and not warnings:
        status = "PASS"
        decision = "PROCEED_TO_STAGE38F_GLD_FEATURE_REVIEW"
    elif daily_rows and feature_rows and joined_rows and not violations:
        status = "WATCH"
        decision = "REVIEW_WARNINGS_BEFORE_STAGE38F_GLD_FEATURE_REVIEW"
    else:
        status = "FAIL"
        decision = "STOP_STAGE38F_GLD_LOADER_FIX_REQUIRED"

    audit = Audit(
        status=status,
        decision=decision,
        workbook_path=str(workbook),
        sheet_name=sheet_name,
        raw_rows_seen=int(meta.get("raw_rows_seen", 0)),
        parsed_rows=len(daily_rows),
        holiday_or_non_numeric_rows_skipped=int(meta.get("holiday_or_non_numeric_rows_skipped", 0)),
        duplicate_date_count=int(meta.get("duplicate_date_count", 0)),
        daily_rows_written=len(daily_rows),
        feature_rows_written=len(feature_rows),
        h1_bars_total=len(h1_bars),
        h1_joined_count=len(joined_rows),
        h1_unjoined_count=len(h1_bars) - len(joined_rows),
        lookahead_violation_count=violations,
        warning_count=len(warnings),
        note_count=len(notes),
        min_observation_date=daily_rows[0]["observation_date"] if daily_rows else None,
        max_observation_date=daily_rows[-1]["observation_date"] if daily_rows else None,
        min_available_from_utc=daily_rows[0]["available_from_utc"] if daily_rows else None,
        max_available_from_utc=daily_rows[-1]["available_from_utc"] if daily_rows else None,
        json_report=json_report,
        md_report=md_report,
    )

    cur = con.cursor()
    cur.execute(
        f'''
        insert into "{AUDIT_TABLE}" (
            created_at_utc, status, decision, workbook_path, sheet_name,
            raw_rows_seen, parsed_rows, holiday_or_non_numeric_rows_skipped,
            duplicate_date_count, daily_rows_written, feature_rows_written,
            h1_bars_total, h1_joined_count, h1_unjoined_count,
            lookahead_violation_count, warning_count, note_count,
            min_observation_date, max_observation_date,
            min_available_from_utc, max_available_from_utc,
            json_report, md_report
        ) values (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        ''',
        (
            utc_now_iso(), audit.status, audit.decision, audit.workbook_path, audit.sheet_name,
            audit.raw_rows_seen, audit.parsed_rows, audit.holiday_or_non_numeric_rows_skipped,
            audit.duplicate_date_count, audit.daily_rows_written, audit.feature_rows_written,
            audit.h1_bars_total, audit.h1_joined_count, audit.h1_unjoined_count,
            audit.lookahead_violation_count, audit.warning_count, audit.note_count,
            audit.min_observation_date, audit.max_observation_date,
            audit.min_available_from_utc, audit.max_available_from_utc,
            audit.json_report, audit.md_report,
        ),
    )
    con.commit()
    write_reports(report_dir, audit, daily_rows, feature_rows)

    print(json.dumps(asdict(audit), indent=2, ensure_ascii=False))
    return 0 if status in {"PASS", "WATCH"} else 2


if __name__ == "__main__":
    raise SystemExit(main())
