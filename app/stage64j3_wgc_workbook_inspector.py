#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import sys
from pathlib import Path
from typing import Any

try:
    import openpyxl
except Exception as exc:  # pragma: no cover
    openpyxl = None
    OPENPYXL_IMPORT_ERROR = str(exc)
else:
    OPENPYXL_IMPORT_ERROR = ""


def utc_now_z() -> str:
    return dt.datetime.now(dt.UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def clean_cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0, 0) else v.isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v)
    s = re.sub(r"\s+", " ", s).strip()
    return s[:240]


def normalize_text(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def row_non_empty(row: list[Any]) -> int:
    return sum(1 for x in row if clean_cell(x) != "")


def detect_header_rows(rows: list[list[Any]], keywords: list[str]) -> list[dict[str, Any]]:
    keys = [normalize_text(k) for k in keywords]
    out = []
    for idx, row in enumerate(rows, start=1):
        cells = [normalize_text(clean_cell(x)) for x in row]
        text = " | ".join(c for c in cells if c)
        if not text:
            continue
        hits = []
        for k in keys:
            if not k:
                continue
            if re.search(rf"\b{re.escape(k)}\b", text):
                hits.append(k)
        if len(set(hits)) >= 2 and row_non_empty(row) >= 2:
            out.append({
                "row_index": idx,
                "keyword_hits": sorted(set(hits)),
                "non_empty_cells": row_non_empty(row),
                "row_preview": [clean_cell(x) for x in row[:16]],
            })
    return out[:10]


def infer_sheet_role(group: str, sheet_name: str, preview_rows: list[list[Any]], config: dict[str, Any]) -> dict[str, Any]:
    role = group
    keywords = config.get("header_keyword_sets", {}).get("etf" if group == "wgc_etf" else "central_bank", [])
    headers = detect_header_rows(preview_rows, keywords)

    text = normalize_text(" ".join([sheet_name] + [" ".join(clean_cell(c) for c in r) for r in preview_rows[:6]]))
    hints = []
    if any(x in text for x in ["flow", "flows", "etf", "aum"]):
        hints.append("ETF_FLOW_OR_HOLDINGS_CANDIDATE")
    if any(x in text for x in ["reserve", "reserves", "country", "ifs"]):
        hints.append("CENTRAL_BANK_RESERVES_CANDIDATE")
    if any(x in text for x in ["change", "changes", "purchase", "purchases", "sale", "sales"]):
        hints.append("CENTRAL_BANK_CHANGES_CANDIDATE")
    if any(x in text for x in ["quarter", "quarterly"]):
        hints.append("QUARTERLY_CANDIDATE")
    if any(x in text for x in ["month", "monthly"]):
        hints.append("MONTHLY_CANDIDATE")

    return {
        "role_hint": ";".join(hints) if hints else "UNKNOWN_FROM_PREVIEW",
        "candidate_header_rows": headers,
    }


def inspect_workbook(path: Path, group: str, config: dict[str, Any]) -> dict[str, Any]:
    if openpyxl is None:
        raise RuntimeError(f"openpyxl import failed: {OPENPYXL_IMPORT_ERROR}")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    max_preview_rows = int(config.get("max_preview_rows", 12))
    max_preview_cols = int(config.get("max_preview_cols", 16))
    max_sheets = int(config.get("max_sheets_per_workbook", 40))

    sheets = []
    for ws in wb.worksheets[:max_sheets]:
        preview_rows: list[list[Any]] = []
        non_empty_preview_rows = 0
        for row in ws.iter_rows(min_row=1, max_row=min(max_preview_rows, ws.max_row or max_preview_rows), max_col=min(max_preview_cols, ws.max_column or max_preview_cols), values_only=True):
            r = list(row)
            preview_rows.append(r)
            if row_non_empty(r):
                non_empty_preview_rows += 1

        role_info = infer_sheet_role(group, ws.title, preview_rows, config)
        sheets.append({
            "sheet_name": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "non_empty_preview_rows": non_empty_preview_rows,
            "role_hint": role_info["role_hint"],
            "candidate_header_rows": role_info["candidate_header_rows"],
            "preview_rows": [[clean_cell(c) for c in r] for r in preview_rows],
        })
    return {
        "group": group,
        "workbook_path": str(path),
        "workbook_name": path.name,
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": sheets,
    }


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in fieldnames})


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=".")
    ap.add_argument("--config", default="configs/stage64j3_wgc_workbook_inspector.json")
    ap.add_argument("--out", default="reports/stage64j3_wgc_workbook_inspector")
    args = ap.parse_args()

    root = Path(args.root).resolve()
    config_path = (root / args.config).resolve()
    out_dir = (root / args.out).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    cfg = load_json(config_path)
    generated_utc = utc_now_z()

    if openpyxl is None:
        summary = {
            "stage": "Stage64J3_WGC_WORKBOOK_INSPECTOR_NO_VALIDATION",
            "status": "BLOCKED_OPENPYXL_NOT_AVAILABLE_NO_PROMOTION",
            "decision": "INSTALL_OR_ENABLE_OPENPYXL_THEN_RERUN_NO_ORDER",
            "openpyxl_error": OPENPYXL_IMPORT_ERROR,
            "promotion": "NO_GO",
            "EA": "NO_GO",
            "paper_order": "NO_GO",
            "paper_live": "NO_GO",
            "live": "NO_GO",
            "validation_allowed_for_order_or_promotion": False,
            "generated_utc": generated_utc,
            "hard_blocks": cfg.get("hard_blocks", []),
        }
        (out_dir / "stage64j3_wgc_workbook_inspector_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
        (out_dir / "stage64j3_wgc_workbook_inspector_report.md").write_text(
            "# Stage64J3 - WGC Workbook Inspector\n\nopenpyxl is not available. No validation or order path is authorized.\n",
            encoding="utf-8",
        )
        return 2

    globs = cfg.get("file_globs", ["*.xlsx", "*.xlsm"])
    workbooks: list[dict[str, Any]] = []
    errors = []
    for group, rel in cfg.get("vendor_roots", {}).items():
        vendor_dir = root / rel
        for pattern in globs:
            for path in sorted(vendor_dir.glob(pattern)):
                try:
                    workbooks.append(inspect_workbook(path, group, cfg))
                except Exception as exc:
                    errors.append({"group": group, "workbook_path": str(path), "error": str(exc)})

    sheet_rows = []
    for wb in workbooks:
        for sh in wb["sheets"]:
            sheet_rows.append({
                "group": wb["group"],
                "workbook_name": wb["workbook_name"],
                "workbook_path": wb["workbook_path"],
                "sheet_name": sh["sheet_name"],
                "max_row": sh["max_row"],
                "max_column": sh["max_column"],
                "role_hint": sh["role_hint"],
                "candidate_header_rows": json.dumps(sh["candidate_header_rows"], ensure_ascii=False),
            })
    write_csv(
        out_dir / "stage64j3_workbook_sheet_inventory.csv",
        sheet_rows,
        ["group", "workbook_name", "workbook_path", "sheet_name", "max_row", "max_column", "role_hint", "candidate_header_rows"],
    )

    summary = {
        "stage": "Stage64J3_WGC_WORKBOOK_INSPECTOR_NO_VALIDATION",
        "status": "WGC_WORKBOOK_INSPECTION_COMPLETE_NO_PROMOTION",
        "decision": "WORKBOOK_STRUCTURE_CAPTURED_BUILD_MAPPER_NEXT_NO_ORDER" if workbooks else "NO_WORKBOOKS_FOUND_CONTINUE_SOURCE_ACQUISITION_OR_PROGRAM_STOP_NO_ORDER",
        "promotion": "NO_GO",
        "EA": "NO_GO",
        "paper_order": "NO_GO",
        "paper_live": "NO_GO",
        "live": "NO_GO",
        "validation_allowed_for_order_or_promotion": False,
        "generated_utc": generated_utc,
        "root": str(root),
        "counts": {
            "workbooks_found": len(workbooks),
            "sheets_inspected": sum(len(wb["sheets"]) for wb in workbooks),
            "inspection_errors": len(errors),
        },
        "workbooks": [
            {
                "group": wb["group"],
                "workbook_name": wb["workbook_name"],
                "workbook_path": wb["workbook_path"],
                "sheet_count": wb["sheet_count"],
                "sheet_names": wb["sheet_names"],
            }
            for wb in workbooks
        ],
        "errors": errors,
        "next_allowed_step": "Stage64J4_WGC_ETF_AND_CENTRAL_BANK_MAPPER_NO_VALIDATION_AFTER_OPERATOR_REVIEW" if workbooks else "ACQUIRE_WGC_WORKBOOKS_OR_PROGRAM_STOP_NO_ORDER",
        "hard_blocks": cfg.get("hard_blocks", []),
        "outputs": {
            "summary_json": str(out_dir / "stage64j3_wgc_workbook_inspector_summary.json"),
            "report_md": str(out_dir / "stage64j3_wgc_workbook_inspector_report.md"),
            "workbook_inventory_json": str(out_dir / "stage64j3_workbook_inventory.json"),
            "sheet_inventory_csv": str(out_dir / "stage64j3_workbook_sheet_inventory.csv"),
        },
    }

    (out_dir / "stage64j3_wgc_workbook_inspector_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    (out_dir / "stage64j3_workbook_inventory.json").write_text(json.dumps(workbooks, indent=2, ensure_ascii=False), encoding="utf-8")

    md = []
    md.append("# Stage64J3 - WGC Workbook Inspector (No Validation)\n")
    md.append(f"Generated UTC: `{generated_utc}`\n")
    md.append("## Status\n")
    md.append(f"- status: `{summary['status']}`\n")
    md.append(f"- decision: `{summary['decision']}`\n")
    md.append("- validation_allowed_for_order_or_promotion: `false`\n")
    md.append("- promotion/paper/live: `NO_GO`\n")
    md.append("\n## Executive conclusion\n")
    md.append("This stage only inspects WGC workbook structures and captures sheet previews/header candidates. It does not map data into strategy inputs, run validation, generate signals, or authorize any order path.\n")
    md.append("\n## Workbooks\n")
    if not workbooks:
        md.append("- none found\n")
    for wb in workbooks:
        md.append(f"\n### `{wb['workbook_name']}` ({wb['group']})\n")
        md.append(f"- path: `{wb['workbook_path']}`\n")
        md.append(f"- sheet_count: `{wb['sheet_count']}`\n")
        for sh in wb["sheets"]:
            md.append(f"\n#### Sheet `{sh['sheet_name']}`\n")
            md.append(f"- rows: `{sh['max_row']}` cols: `{sh['max_column']}`\n")
            md.append(f"- role_hint: `{sh['role_hint']}`\n")
            if sh["candidate_header_rows"]:
                md.append("- candidate headers:\n")
                for h in sh["candidate_header_rows"][:4]:
                    md.append(f"  - row {h['row_index']}: hits={h['keyword_hits']} preview={h['row_preview']}\n")
            md.append("- preview:\n\n")
            md.append("```text\n")
            for r in sh["preview_rows"][: int(cfg.get("max_preview_rows", 12))]:
                md.append(" | ".join(r) + "\n")
            md.append("```\n")
    if errors:
        md.append("\n## Errors\n")
        for e in errors:
            md.append(f"- `{e.get('workbook_path')}`: {e.get('error')}\n")
    md.append("\n## Next allowed step\n")
    md.append(f"`{summary['next_allowed_step']}`\n")
    (out_dir / "stage64j3_wgc_workbook_inspector_report.md").write_text("".join(md), encoding="utf-8")

    return 0 if workbooks else 1


if __name__ == "__main__":
    raise SystemExit(main())
