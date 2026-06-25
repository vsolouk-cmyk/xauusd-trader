#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import math
import re
from calendar import monthrange
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

try:
    import openpyxl
except Exception as exc:  # pragma: no cover
    openpyxl = None
    OPENPYXL_IMPORT_ERROR = str(exc)
else:
    OPENPYXL_IMPORT_ERROR = ""


def utc_now() -> str:
    return dt.datetime.now(dt.UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Stage64J4 WGC ETF and central-bank streaming mapper - no validation/no order")
    p.add_argument("--root", default=".", help="Repo root")
    p.add_argument("--config", default="configs/stage64j4_wgc_etf_central_bank_mapper.json")
    p.add_argument("--out", default="reports/stage64j4_wgc_etf_central_bank_mapper")
    return p.parse_args()


def read_json(path: Path) -> Dict[str, Any]:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def write_json(path: Path, obj: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        json.dump(obj, f, indent=2, ensure_ascii=False)


def write_csv(path: Path, rows: List[Dict[str, Any]], fieldnames: List[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in fieldnames})


def as_text(x: Any) -> str:
    if x is None:
        return ""
    if isinstance(x, (dt.datetime, dt.date)):
        return x.isoformat()
    return str(x).strip()


def to_date(x: Any) -> Optional[dt.date]:
    if x is None or x == "":
        return None
    if isinstance(x, dt.datetime):
        return x.date()
    if isinstance(x, dt.date):
        return x
    s = str(x).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(s[:10], fmt).date()
        except Exception:
            pass
    try:
        return dt.date.fromisoformat(s[:10])
    except Exception:
        return None


def to_float(x: Any) -> Optional[float]:
    if x is None or x == "":
        return None
    if isinstance(x, bool):
        return None
    if isinstance(x, (int, float)):
        v = float(x)
        if math.isnan(v) or math.isinf(v):
            return None
        return v
    s = str(x).strip().replace(",", "")
    if not s:
        return None
    try:
        v = float(s)
    except Exception:
        return None
    if math.isnan(v) or math.isinf(v):
        return None
    return v


def iso_date(d: dt.date) -> str:
    return d.isoformat()


def iso_utc(d: dt.date) -> str:
    return d.isoformat() + "T00:00:00Z"


def month_end(d: dt.date) -> dt.date:
    return dt.date(d.year, d.month, monthrange(d.year, d.month)[1])


def normalize_id(*parts: str) -> str:
    joined = "_".join(p for p in parts if p)
    joined = re.sub(r"[^A-Za-z0-9]+", "_", joined).strip("_")
    return joined.upper()[:180] or "UNKNOWN"


def find_latest(globs: Iterable[str], root: Path) -> Optional[Path]:
    files: List[Path] = []
    for g in globs:
        files.extend(root.glob(g))
    files = [p for p in files if p.is_file()]
    if not files:
        return None
    return sorted(files, key=lambda p: p.stat().st_mtime, reverse=True)[0]


def load_sheet_rows(path: Path, sheet_name: str) -> Tuple[List[Tuple[Any, ...]], List[str]]:
    if openpyxl is None:
        raise RuntimeError(f"openpyxl unavailable: {OPENPYXL_IMPORT_ERROR}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise RuntimeError(f"sheet not found: {sheet_name}; available={wb.sheetnames}")
        ws = wb[sheet_name]
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        sheet_names = list(wb.sheetnames)
        return rows, sheet_names
    finally:
        try:
            wb.close()
        except Exception:
            pass


def get_cell(rows: List[Tuple[Any, ...]], row_1: int, col_1: int) -> Any:
    r = row_1 - 1
    c = col_1 - 1
    if r < 0 or c < 0 or r >= len(rows):
        return None
    row = rows[r]
    if c >= len(row):
        return None
    return row[c]


def find_header_row(rows: List[Tuple[Any, ...]], first_cell: str = "Date", max_rows: int = 25) -> Optional[int]:
    needle = first_cell.strip().lower()
    for idx, row in enumerate(rows[:max_rows], start=1):
        val = as_text(row[0] if row else "").strip().lower()
        if val == needle:
            return idx
    return None


def find_central_bank_header_row(rows: List[Tuple[Any, ...]], max_rows: int = 25) -> Optional[int]:
    for idx, row in enumerate(rows[:max_rows], start=1):
        vals = [as_text(x).lower() for x in row[:8]]
        if "country lookup column" in vals and "country" in vals and "comments" in vals:
            return idx
    return None


def map_etf(config: Dict[str, Any], root: Path) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    etf_cfg = config["etf"]
    path = find_latest(etf_cfg["vendor_globs"], root)
    if path is None:
        return [], {"ok": False, "issue": "ETF workbook missing", "rows": 0}

    sheet_name = etf_cfg.get("sheet", "Demand by month")
    rows_data, sheet_names = load_sheet_rows(path, sheet_name)

    header_row = find_header_row(rows_data, "Date")
    if not header_row:
        return [], {"ok": False, "issue": "ETF Date header row not found", "rows": 0, "workbook": str(path), "sheet": sheet_name, "available_sheets": sheet_names}

    ticker_row = max(1, header_row - 5)
    fund_type_row = max(1, header_row - 3)
    region_row = max(1, header_row - 2)
    country_row = max(1, header_row - 1)
    data_start_row = header_row + 1
    fund_start_col = int(etf_cfg.get("fund_start_column_1_indexed", 6))
    release_lag_days = int(etf_cfg.get("release_lag_days_after_period_end", 10))
    available_lag_days = int(etf_cfg.get("available_lag_days_after_period_end", release_lag_days + 1))

    max_cols = max((len(r) for r in rows_data), default=0)
    column_meta: Dict[int, Dict[str, str]] = {}
    for c in range(fund_start_col, max_cols + 1):
        fund_name = as_text(get_cell(rows_data, header_row, c))
        ticker = as_text(get_cell(rows_data, ticker_row, c))
        fund_type = as_text(get_cell(rows_data, fund_type_row, c))
        region = as_text(get_cell(rows_data, region_row, c))
        country = as_text(get_cell(rows_data, country_row, c))
        if not fund_name and not ticker:
            continue
        etf_id = normalize_id(ticker or fund_name)
        column_meta[c] = {
            "etf_id": etf_id,
            "ticker": ticker,
            "fund_name": fund_name,
            "fund_type": fund_type,
            "region": region,
            "country": country,
        }

    out_rows: List[Dict[str, Any]] = []
    for r_1 in range(data_start_row, len(rows_data) + 1):
        d = to_date(get_cell(rows_data, r_1, 1))
        if d is None:
            continue
        release_date = d + dt.timedelta(days=release_lag_days)
        available_after = d + dt.timedelta(days=available_lag_days)
        for c, meta in column_meta.items():
            v = to_float(get_cell(rows_data, r_1, c))
            if v is None:
                continue
            out_rows.append({
                "date_utc": iso_date(d),
                "etf_id": meta["etf_id"],
                "holdings_tonnes_or_flow": f"{v:.12g}",
                "source": "WGC_ETF_DEMAND_BY_MONTH_TONNES_ASSUMED_MONTHLY_RELEASE_LAG|"
                          f"workbook={path.name}|sheet={sheet_name}|ticker={meta['ticker']}|fund={meta['fund_name']}|region={meta['region']}|country={meta['country']}",
                "release_time_utc": iso_utc(release_date),
                "available_after_utc": iso_utc(available_after),
            })

    out_rows.sort(key=lambda x: (x["date_utc"], x["etf_id"]))
    diag = {
        "ok": bool(out_rows),
        "workbook": str(path),
        "sheet": sheet_name,
        "header_row": header_row,
        "fund_columns": len(column_meta),
        "rows": len(out_rows),
        "first_date": out_rows[0]["date_utc"] if out_rows else None,
        "last_date": out_rows[-1]["date_utc"] if out_rows else None,
        "release_lag_days_after_period_end": release_lag_days,
        "available_lag_days_after_period_end": available_lag_days,
        "unit_interpretation": "monthly ETF demand/flow in tonnes by fund from WGC Demand by month sheet",
        "loader_mode": "streamed_iter_rows_then_in_memory_sheet_matrix",
    }
    return out_rows, diag


def map_central_bank(config: Dict[str, Any], root: Path) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    cb_cfg = config["central_bank"]
    path = find_latest(cb_cfg["vendor_globs"], root)
    if path is None:
        return [], {"ok": False, "issue": "Central-bank changes workbook missing", "rows": 0}

    sheet_name = cb_cfg.get("sheet", "Monthly")
    rows_data, sheet_names = load_sheet_rows(path, sheet_name)

    header_row = find_central_bank_header_row(rows_data)
    if not header_row:
        return [], {"ok": False, "issue": "Central-bank header row not found", "rows": 0, "workbook": str(path), "sheet": sheet_name, "available_sheets": sheet_names}

    start_col = int(cb_cfg.get("first_period_column_1_indexed", 4))
    release_lag_days = int(cb_cfg.get("release_lag_days_after_period_end", 75))
    max_cols = max((len(r) for r in rows_data), default=0)

    period_dates: Dict[int, dt.date] = {}
    for c in range(start_col, max_cols + 1):
        d = to_date(get_cell(rows_data, header_row, c))
        if d is None:
            continue
        if d < dt.date(2002, 1, 1):
            continue
        period_dates[c] = d

    by_period: Dict[dt.date, Dict[str, Any]] = {d: {"sum": 0.0, "non_null_count": 0, "non_zero_count": 0} for d in period_dates.values()}

    for r_1 in range(header_row + 1, len(rows_data) + 1):
        country = as_text(get_cell(rows_data, r_1, 2))
        if not country or country.lower() in {"country", "comments"}:
            continue
        for c, period_start in period_dates.items():
            v = to_float(get_cell(rows_data, r_1, c))
            if v is None:
                continue
            rec = by_period[period_start]
            rec["sum"] += v
            rec["non_null_count"] += 1
            if abs(v) > 0:
                rec["non_zero_count"] += 1

    out_rows: List[Dict[str, Any]] = []
    for period_start in sorted(by_period):
        rec = by_period[period_start]
        if rec["non_null_count"] <= 0:
            continue
        pend = month_end(period_start)
        release_date = pend + dt.timedelta(days=release_lag_days)
        out_rows.append({
            "period_start": iso_date(period_start),
            "period_end": iso_date(pend),
            "demand_value": f"{rec['sum']:.12g}",
            "unit": "tonnes_net_change_sum_all_reported_countries",
            "source": "WGC_CHANGES_LATEST_MONTHLY_AGGREGATED_ALL_COUNTRIES_ASSUMED_75D_RELEASE_LAG|"
                      f"workbook={path.name}|sheet={sheet_name}|non_null_country_values={rec['non_null_count']}|non_zero_country_values={rec['non_zero_count']}",
            "release_date_utc": iso_utc(release_date),
            "available_after_utc": iso_utc(release_date),
        })

    diag = {
        "ok": bool(out_rows),
        "workbook": str(path),
        "sheet": sheet_name,
        "header_row": header_row,
        "period_columns": len(period_dates),
        "rows": len(out_rows),
        "first_period_start": out_rows[0]["period_start"] if out_rows else None,
        "last_period_start": out_rows[-1]["period_start"] if out_rows else None,
        "release_lag_days_after_period_end": release_lag_days,
        "aggregation": "sum monthly reported country reserve changes into global central-bank net demand proxy",
        "loader_mode": "streamed_iter_rows_then_in_memory_sheet_matrix",
    }
    return out_rows, diag


def main() -> int:
    args = parse_args()
    root = Path(args.root).resolve()
    config_path = (root / args.config).resolve()
    out = (root / args.out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    config = read_json(config_path)

    generated = utc_now()
    hard_blocks = [
        "NO_PAPER_ORDER",
        "NO_EA_PROMOTION",
        "NO_PAPER_LIVE",
        "NO_LIVE",
        "NO_BROKER_CONNECTION",
        "NO_VALIDATION_SCAN_IN_STAGE64J4",
        "NO_REDUCED_SCOPE_RETEST",
        "NO_RESCUE_FILTERING",
        "NO_NEW_INTRADAY_SCAN",
    ]

    if openpyxl is None:
        summary = {
            "stage": "Stage64J4_WGC_ETF_AND_CENTRAL_BANK_MAPPER_NO_VALIDATION_LOADERFIX1",
            "status": "MAPPER_BLOCKED_OPENPYXL_MISSING_NO_PROMOTION",
            "decision": "INSTALL_OPENPYXL_AND_RERUN_NO_ORDER",
            "generated_utc": generated,
            "validation_allowed_for_order_or_promotion": False,
            "openpyxl_error": OPENPYXL_IMPORT_ERROR,
            "hard_blocks": hard_blocks,
        }
        write_json(out / "stage64j4_wgc_etf_central_bank_mapper_summary.json", summary)
        return 2

    etf_rows, etf_diag = map_etf(config, root)
    cb_rows, cb_diag = map_central_bank(config, root)

    etf_target = root / config["outputs"]["etf_target_file"]
    cb_target = root / config["outputs"]["central_bank_target_file"]

    if etf_rows:
        write_csv(etf_target, etf_rows, ["date_utc", "etf_id", "holdings_tonnes_or_flow", "source", "release_time_utc", "available_after_utc"])
    if cb_rows:
        write_csv(cb_target, cb_rows, ["period_start", "period_end", "demand_value", "unit", "source", "release_date_utc", "available_after_utc"])

    diagnostics_rows = [
        {"source": "ETF", **etf_diag},
        {"source": "CENTRAL_BANK", **cb_diag},
    ]
    write_csv(out / "stage64j4_mapper_diagnostics.csv", diagnostics_rows, sorted({k for r in diagnostics_rows for k in r.keys()}))

    etf_min = int(config["etf"].get("minimum_rows_for_j1_preflight", 1000))
    cb_min = int(config["central_bank"].get("minimum_rows_for_j1_preflight", 40))
    etf_ok = len(etf_rows) >= etf_min
    cb_ok = len(cb_rows) >= cb_min
    mapper_ok = etf_ok and cb_ok

    decision = (
        "WGC_ETF_AND_CENTRAL_BANK_MAPPED_RERUN_STAGE64J1_EVENT_CALENDAR_OR_BROKER_STILL_BLOCKED_NO_ORDER"
        if mapper_ok else
        "WGC_MAPPING_INCOMPLETE_REVIEW_WORKBOOK_STRUCTURE_NO_ORDER"
    )
    status = "WGC_MAPPING_COMPLETE_NO_PROMOTION" if mapper_ok else "WGC_MAPPING_INCOMPLETE_NO_PROMOTION"

    summary = {
        "stage": "Stage64J4_WGC_ETF_AND_CENTRAL_BANK_MAPPER_NO_VALIDATION_LOADERFIX1",
        "status": status,
        "decision": decision,
        "promotion": "NO_GO",
        "EA": "NO_GO",
        "paper_order": "NO_GO",
        "paper_live": "NO_GO",
        "live": "NO_GO",
        "validation_allowed_for_order_or_promotion": False,
        "generated_utc": generated,
        "root": str(root),
        "inputs": {"config": str(config_path)},
        "outputs": {
            "etf_target_file": str(etf_target),
            "central_bank_target_file": str(cb_target),
            "summary_json": str(out / "stage64j4_wgc_etf_central_bank_mapper_summary.json"),
            "report_md": str(out / "stage64j4_wgc_etf_central_bank_mapper_report.md"),
            "diagnostics_csv": str(out / "stage64j4_mapper_diagnostics.csv"),
        },
        "mapping_results": {
            "etf": {**etf_diag, "row_count_ok": etf_ok, "minimum_rows_for_j1_preflight": etf_min, "target_file": str(etf_target)},
            "central_bank": {**cb_diag, "row_count_ok": cb_ok, "minimum_rows_for_j1_preflight": cb_min, "target_file": str(cb_target)},
        },
        "remaining_blockers_expected_after_mapping": [
            "SRC_MACRO_EVENT_CALENDAR_ARCHIVE unless real historical archive is acquired or J1 accepts forward-only governance",
            "SRC_BROKER_OR_SPOT_GOLD_D1_ALIGNMENT before commercialization claim",
        ],
        "source_lag_assumptions": [
            "ETF monthly demand rows use a conservative assumed release/availability lag after period-end because exact historical release timestamps are not present in workbook preview.",
            "Central-bank monthly net demand proxy uses WGC monthly reserve changes aggregated across countries with a conservative 75-day official-release lag.",
            "These files are source-completion inputs only and do not authorize validation or order paths.",
        ],
        "next_allowed_step": "RERUN_STAGE64J1_SOURCE_PREFLIGHT_OR_PROGRAM_STOP_NO_ORDER",
        "hard_blocks": hard_blocks,
    }

    write_json(out / "stage64j4_wgc_etf_central_bank_mapper_summary.json", summary)

    report = out / "stage64j4_wgc_etf_central_bank_mapper_report.md"
    with report.open("w", encoding="utf-8") as f:
        f.write("# Stage64J4 - WGC ETF and Central-Bank Mapper LoaderFix1 (No Validation)\n\n")
        f.write(f"Generated UTC: `{generated}`\n\n")
        f.write("## Status\n\n")
        f.write(f"- status: `{status}`\n")
        f.write(f"- decision: `{decision}`\n")
        f.write("- validation_allowed_for_order_or_promotion: `false`\n")
        f.write("- promotion/paper/live: `NO_GO`\n\n")
        f.write("## Executive conclusion\n\n")
        f.write("Stage64J4 LoaderFix1 maps WGC ETF and central-bank workbooks into the Stage64 raw source schemas using a streaming worksheet reader. It does not run validation, generate signals, connect to a broker, or authorize any order path.\n\n")
        f.write("## Mapping results\n\n")
        f.write("| source | ok | rows | first | last | target |\n")
        f.write("|---|---:|---:|---|---|---|\n")
        f.write(f"| ETF | {etf_ok} | {len(etf_rows)} | {etf_diag.get('first_date')} | {etf_diag.get('last_date')} | `{etf_target.relative_to(root)}` |\n")
        f.write(f"| Central bank | {cb_ok} | {len(cb_rows)} | {cb_diag.get('first_period_start')} | {cb_diag.get('last_period_start')} | `{cb_target.relative_to(root)}` |\n\n")
        f.write("## Loader fix\n\n")
        f.write("The previous mapper could be slow on read-only workbooks because it repeatedly accessed worksheet cells. LoaderFix1 streams each relevant sheet once with `iter_rows(values_only=True)` and then maps the small in-memory matrix.\n\n")
        f.write("## Source interpretation\n\n")
        f.write("- ETF source: WGC `Demand by month` sheet, long-form by fund, monthly tonnes.\n")
        f.write("- Central-bank source: WGC `Changes_latest` `Monthly` sheet, aggregated monthly net reserve changes across reported countries.\n")
        f.write("- Release/availability timestamps are conservative lag assumptions because exact historical publication timestamps are not embedded in these workbooks.\n\n")
        f.write("## Remaining blockers\n\n")
        f.write("- Historical event calendar remains blocked unless a real archive is acquired, or governance formally keeps it forward-only.\n")
        f.write("- Broker/spot D1 alignment remains blocked before any commercialization or broker XAUUSD claim.\n\n")
        f.write("## Operational decision\n\n")
        f.write("No validation scan, signal generation, paper-order, paper-live, live, EA promotion, or broker connection is authorized by Stage64J4.\n\n")
        f.write("## Next allowed step\n\n")
        f.write("`RERUN_STAGE64J1_SOURCE_PREFLIGHT_OR_PROGRAM_STOP_NO_ORDER`\n")

    return 0 if mapper_ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
