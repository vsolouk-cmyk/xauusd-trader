#!/usr/bin/env python3
"""Inventory local XAUUSD macro/regime data without network or execution access."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sqlite3
import tempfile
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping

PROGRAM = "XAUUSD_MACRO_DATA_READINESS_INVENTORY_V1_LOCAL_ONLY"
UTC = timezone.utc

SERIES_DEFINITIONS = {
    "usd_broad_daily": {
        "required": True,
        "identifiers": ["DTWEXBGS", "broad dollar", "trade weighted", "trade-weighted"],
        "filename_tokens": ["dtwexbgs", "broad_dollar", "trade_weighted", "trade-weighted"],
        "quality_note": "Prefer Federal Reserve broad trade-weighted USD index. ICE DXY is not required.",
    },
    "real_yield_10y_daily": {
        "required": True,
        "identifiers": ["DFII10", "10 year real", "10-year real", "real yield", "TIPS"],
        "filename_tokens": ["dfii10", "real_yield", "tips", "treasury_real"],
        "quality_note": "Prefer 10-year real Treasury yield with publication-safe availability.",
    },
    "nominal_yield_10y_daily": {
        "required": True,
        "identifiers": ["DGS10", "10 year treasury", "10-year treasury", "nominal yield"],
        "filename_tokens": ["dgs10", "nominal_yield", "treasury_10y"],
        "quality_note": "Daily 10-year nominal Treasury yield.",
    },
    "breakeven_10y_daily": {
        "required": True,
        "identifiers": ["T10YIE", "10 year breakeven", "10-year breakeven", "breakeven"],
        "filename_tokens": ["t10yie", "breakeven"],
        "quality_note": "Daily 10-year inflation breakeven.",
    },
    "vix_daily": {
        "required": True,
        "identifiers": ["VIXCLS", "VIX Index", "VIX"],
        "filename_tokens": ["vix_history", "vixcls", "vix_daily"],
        "quality_note": "Daily VIX close; must be shifted to next-session availability.",
    },
    "gvz_daily": {
        "required": True,
        "identifiers": ["GVZ", "Gold ETF Volatility"],
        "filename_tokens": ["gvz_history", "gvz_daily"],
        "quality_note": "Daily Cboe Gold ETF Volatility Index.",
    },
    "cftc_gold_weekly": {
        "required": True,
        "identifiers": ["GOLD - COMMODITY EXCHANGE INC", "COMMODITY EXCHANGE INC", "CFTC", "COT"],
        "filename_tokens": ["fut_disagg", "cftc", "cot", "commitment"],
        "quality_note": "Weekly gold positioning; availability must use release date, not report date.",
    },
    "official_event_calendar": {
        "required": True,
        "identifiers": ["event_time_utc", "BLS", "BEA", "FED", "FOMC"],
        "filename_tokens": ["official_event", "event_calendar", "historical_event_context"],
        "quality_note": "Existing official event calendar and blackout context.",
    },
    "gold_etf_holdings_or_flows": {
        "required": False,
        "identifiers": ["GLD", "gold ETF", "ETF holdings", "ETF flows", "SPDR Gold"],
        "filename_tokens": ["gld", "gold_etf", "etf_holdings", "etf_flows"],
        "quality_note": "Desirable daily investment-flow proxy; source and publication timing must be documented.",
    },
    "wgc_quarterly_demand": {
        "required": False,
        "identifiers": ["World Gold Council", "Gold Demand Trends", "central banks", "central bank demand"],
        "filename_tokens": ["wgc", "gold_demand", "gdt", "central_bank"],
        "quality_note": "Quarterly regime prior only; use publication date, never quarter-end hindsight.",
    },
}

SUPPORTED_SUFFIXES = {".csv", ".tsv", ".txt", ".json", ".sqlite", ".sqlite3", ".db", ".xlsx", ".xlsm"}
DATE_NAME_RE = re.compile(r"(date|time|timestamp|utc|period|observation)", re.I)


class InventoryError(RuntimeError):
    pass


def iso_utc() -> str:
    return datetime.now(tz=UTC).isoformat().replace("+00:00", "Z")


def atomic_write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, delete=False) as handle:
        handle.write(text)
        tmp = handle.name
    os.replace(tmp, path)


def write_json(path: Path, payload: Any) -> None:
    atomic_write(path, json.dumps(payload, indent=2, ensure_ascii=False, sort_keys=True) + "\n")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def expand_path(root: Path, value: str) -> Path:
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = root / path
    return path.resolve()


def safe_file_iter(search_root: Path, max_files: int) -> Iterable[Path]:
    if not search_root.exists():
        return
    count = 0
    for path in sorted(search_root.rglob("*")):
        if count >= max_files:
            break
        try:
            if path.is_symlink() or not path.is_file():
                continue
        except OSError:
            continue
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        if any(part in {".git", "_archive", "__pycache__", "node_modules"} for part in path.parts):
            continue
        count += 1
        yield path


def read_text_sample(path: Path, max_bytes: int = 256_000) -> str:
    with path.open("rb") as handle:
        raw = handle.read(max_bytes)
    return raw.decode("utf-8-sig", errors="replace")


def sniff_delimiter(text: str, suffix: str) -> str:
    if suffix == ".tsv":
        return "\t"
    try:
        return csv.Sniffer().sniff(text[:8192], delimiters=",\t;|").delimiter
    except csv.Error:
        return "\t" if text.count("\t") > text.count(",") else ","


def summarize_delimited(path: Path) -> dict[str, Any]:
    text = read_text_sample(path)
    delimiter = sniff_delimiter(text, path.suffix.lower())
    lines = text.splitlines()
    rows = list(csv.reader(lines[:1001], delimiter=delimiter))
    headers = [str(x).strip() for x in (rows[0] if rows else [])]
    sample_rows = rows[1:] if len(rows) > 1 else []
    first_values = [" | ".join(map(str, row[:12])) for row in sample_rows[:20]]
    date_columns = [name for name in headers if DATE_NAME_RE.search(name)]
    return {
        "kind": "delimited",
        "delimiter": "\\t" if delimiter == "\t" else delimiter,
        "headers": headers,
        "sample_row_count": len(sample_rows),
        "date_columns": date_columns,
        "text_evidence": "\n".join([",".join(headers)] + first_values)[:50_000],
    }


def summarize_json(path: Path) -> dict[str, Any]:
    text = read_text_sample(path)
    try:
        payload = json.loads(text)
        if isinstance(payload, dict):
            keys = list(payload.keys())[:100]
            evidence = json.dumps({k: payload[k] for k in keys[:20]}, ensure_ascii=False, default=str)
        elif isinstance(payload, list):
            keys = []
            evidence = json.dumps(payload[:10], ensure_ascii=False, default=str)
        else:
            keys = []
            evidence = str(payload)
        return {"kind": "json", "keys": keys, "text_evidence": evidence[:50_000]}
    except Exception as exc:
        return {"kind": "json_unparsed", "error": f"{type(exc).__name__}: {exc}", "text_evidence": text[:50_000]}


def summarize_sqlite(path: Path) -> dict[str, Any]:
    uri = f"file:{path}?mode=ro"
    conn = sqlite3.connect(uri, uri=True)
    tables: dict[str, Any] = {}
    evidence_parts: list[str] = []
    try:
        table_names = [
            row[0] for row in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"
            )
        ]
        for table in table_names[:100]:
            cols = [row[1] for row in conn.execute(f'PRAGMA table_info("{table}")')]
            try:
                count = int(conn.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0])
            except Exception:
                count = None
            date_cols = [c for c in cols if DATE_NAME_RE.search(c)]
            minmax = {}
            for col in date_cols[:3]:
                try:
                    lo, hi = conn.execute(
                        f'SELECT MIN("{col}"), MAX("{col}") FROM "{table}"'
                    ).fetchone()
                    minmax[col] = {"min": lo, "max": hi}
                except Exception:
                    pass
            tables[table] = {
                "columns": cols,
                "row_count": count,
                "date_ranges": minmax,
            }
            evidence_parts.append(f"{table} {' '.join(cols)}")
    finally:
        conn.close()
    return {"kind": "sqlite", "tables": tables, "text_evidence": "\n".join(evidence_parts)[:50_000]}


def summarize_xlsx(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        return {
            "kind": "xlsx_unread",
            "error": f"openpyxl unavailable: {exc}",
            "text_evidence": path.name,
        }
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, Any] = {}
    evidence_parts: list[str] = []
    try:
        for ws in wb.worksheets[:30]:
            rows = []
            for row in ws.iter_rows(min_row=1, max_row=25, values_only=True):
                rows.append(["" if v is None else str(v) for v in row[:30]])
            headers = rows[0] if rows else []
            evidence = "\n".join(" | ".join(r) for r in rows)
            sheets[ws.title] = {"headers": headers, "sample_rows": len(rows)}
            evidence_parts.append(ws.title + "\n" + evidence)
    finally:
        wb.close()
    return {"kind": "xlsx", "sheets": sheets, "text_evidence": "\n".join(evidence_parts)[:80_000]}


def summarize_file(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        return summarize_delimited(path)
    if suffix == ".json":
        return summarize_json(path)
    if suffix in {".sqlite", ".sqlite3", ".db"}:
        return summarize_sqlite(path)
    if suffix in {".xlsx", ".xlsm"}:
        return summarize_xlsx(path)
    raise InventoryError(f"unsupported file: {path}")


def classify(path: Path, summary: Mapping[str, Any]) -> list[dict[str, Any]]:
    name = path.name.lower()
    evidence = (name + "\n" + str(summary.get("text_evidence", ""))).lower()
    results: list[dict[str, Any]] = []
    for series, spec in SERIES_DEFINITIONS.items():
        hits = []
        for token in spec["filename_tokens"]:
            if token.lower() in name:
                hits.append(f"filename:{token}")
        for token in spec["identifiers"]:
            if token.lower() in evidence:
                hits.append(f"content:{token}")
        if not hits:
            continue
        score = min(1.0, 0.25 * len(set(hits)) + (0.25 if any(h.startswith("filename:") for h in hits) else 0.0))
        results.append({
            "series": series,
            "confidence": round(score, 3),
            "evidence": sorted(set(hits)),
        })
    return sorted(results, key=lambda x: (-x["confidence"], x["series"]))


def run_inventory(root: Path, config_path: Path, out_dir: Path) -> dict[str, Any]:
    config = json.loads(config_path.read_text(encoding="utf-8"))
    search_roots = [expand_path(root, value) for value in config["search_roots"]]
    max_files = int(config.get("max_files_per_root", 5000))
    max_bytes = int(config.get("max_file_size_bytes", 800_000_000))
    out_dir.mkdir(parents=True, exist_ok=True)

    file_rows: list[dict[str, Any]] = []
    candidates_by_series: dict[str, list[dict[str, Any]]] = defaultdict(list)
    errors: list[dict[str, str]] = []
    seen: set[Path] = set()

    for search_root in search_roots:
        for path in safe_file_iter(search_root, max_files):
            resolved = path.resolve()
            if resolved in seen:
                continue
            seen.add(resolved)
            try:
                size = path.stat().st_size
                if size > max_bytes:
                    file_rows.append({
                        "path": str(path),
                        "size": size,
                        "status": "SKIPPED_SIZE_LIMIT",
                        "classifications": [],
                    })
                    continue
                summary = summarize_file(path)
                classes = classify(path, summary)
                record = {
                    "path": str(path),
                    "size": size,
                    "sha256": sha256_file(path),
                    "status": "INSPECTED",
                    "summary": {k: v for k, v in summary.items() if k != "text_evidence"},
                    "classifications": classes,
                }
                file_rows.append(record)
                for item in classes:
                    candidates_by_series[item["series"]].append({
                        "path": str(path),
                        "confidence": item["confidence"],
                        "evidence": item["evidence"],
                        "size": size,
                        "sha256": record["sha256"],
                        "kind": summary.get("kind"),
                        "summary": record["summary"],
                    })
            except Exception as exc:
                errors.append({"path": str(path), "error": f"{type(exc).__name__}: {exc}"})

    readiness: dict[str, Any] = {}
    required_missing = []
    for series, spec in SERIES_DEFINITIONS.items():
        candidates = sorted(
            candidates_by_series.get(series, []),
            key=lambda x: (-x["confidence"], x["path"]),
        )
        strong = [c for c in candidates if c["confidence"] >= 0.5]
        present = bool(strong)
        readiness[series] = {
            "required": bool(spec["required"]),
            "present": present,
            "candidate_count": len(candidates),
            "strong_candidate_count": len(strong),
            "quality_note": spec["quality_note"],
            "top_candidates": candidates[:10],
        }
        if spec["required"] and not present:
            required_missing.append(series)

    decision = (
        "PASS_CORE_MACRO_DATA_READY_FOR_CAUSAL_PANEL_BUILD"
        if not required_missing
        else "PARTIAL_CORE_MACRO_DATA_REQUIRES_DOWNLOADS"
    )
    summary = {
        "program": PROGRAM,
        "generated_utc": iso_utc(),
        "decision": decision,
        "pass": not required_missing,
        "paper_order_allowed": False,
        "demo_order_allowed": False,
        "live_order_allowed": False,
        "search_roots": [str(p) for p in search_roots],
        "files_inspected": sum(1 for r in file_rows if r["status"] == "INSPECTED"),
        "files_skipped_size": sum(1 for r in file_rows if r["status"] == "SKIPPED_SIZE_LIMIT"),
        "inspection_errors": len(errors),
        "required_missing": required_missing,
        "readiness": readiness,
        "required_next_action": (
            "BUILD_CAUSAL_MACRO_PANEL_FROM_EXISTING_DATA"
            if not required_missing
            else "DOWNLOAD_ONLY_MISSING_OFFICIAL_SERIES_THEN_BUILD_CAUSAL_MACRO_PANEL"
        ),
    }
    write_json(out_dir / "macro_data_readiness_summary.json", summary)
    write_json(out_dir / "macro_data_inventory_files.json", file_rows)
    write_json(out_dir / "macro_data_inventory_errors.json", errors)
    write_json(out_dir / "macro_data_source_contract.json", {
        "program": PROGRAM,
        "generated_utc": summary["generated_utc"],
        "series_definitions": SERIES_DEFINITIONS,
        "availability_rules": {
            "daily_close_series": "Use from next trading session; never same-day intraday.",
            "cftc_weekly": "Use release availability, not Tuesday report date; conservative Monday availability is acceptable.",
            "quarterly_wgc": "Use publication date, never quarter-end.",
            "scheduled_events": "Use scheduled release timestamps and blackout windows only.",
        },
        "execution_allowed": False,
    })
    return summary


def collect(root: Path, out_dir: Path, download_dir: Path) -> Path:
    required = [
        out_dir / "macro_data_readiness_summary.json",
        out_dir / "macro_data_inventory_files.json",
        out_dir / "macro_data_inventory_errors.json",
        out_dir / "macro_data_source_contract.json",
    ]
    missing = [str(p) for p in required if not p.is_file()]
    if missing:
        raise InventoryError("inventory outputs missing: " + ", ".join(missing))
    download_dir.mkdir(parents=True, exist_ok=True)
    output = download_dir / "XAUUSD_MACRO_DATA_READINESS_RESULTS.zip"
    manifest = {"program": PROGRAM, "generated_utc": iso_utc(), "files": []}
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in required:
            data = path.read_bytes()
            arc = f"reports/{path.name}"
            archive.writestr(arc, data)
            manifest["files"].append({
                "archive_path": arc,
                "sha256": hashlib.sha256(data).hexdigest(),
                "size": len(data),
            })
        archive.writestr("RESULTS_MANIFEST.json", json.dumps(manifest, indent=2) + "\n")
    return output


def resolve_paths(args: argparse.Namespace) -> tuple[Path, Path, Path, Path]:
    root = Path(args.root).expanduser().resolve()
    config = expand_path(root, args.config)
    out = expand_path(root, args.out)
    download = Path(args.download_dir).expanduser().resolve()
    return root, config, out, download


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["preflight", "run", "collect"])
    parser.add_argument("--root", default=".")
    parser.add_argument("--config", default="config/xauusd_macro_data_inventory_v1.json")
    parser.add_argument("--out", default="reports/xauusd_macro_data_readiness")
    parser.add_argument("--download-dir", default="~/Downloads")
    args = parser.parse_args()
    root, config, out, download = resolve_paths(args)
    try:
        if args.command == "preflight":
            if not config.is_file():
                raise InventoryError(f"config missing: {config}")
            payload = {
                "program": PROGRAM,
                "generated_utc": iso_utc(),
                "decision": "PASS_MACRO_DATA_INVENTORY_PREFLIGHT_LOCAL_ONLY",
                "pass": True,
                "paper_order_allowed": False,
                "demo_order_allowed": False,
                "live_order_allowed": False,
                "config": str(config),
                "out": str(out),
            }
        elif args.command == "run":
            payload = run_inventory(root, config, out)
        else:
            output = collect(root, out, download)
            payload = {
                "program": PROGRAM,
                "generated_utc": iso_utc(),
                "decision": "PASS_MACRO_DATA_READINESS_RESULTS_PACK_CREATED",
                "pass": True,
                "paper_order_allowed": False,
                "demo_order_allowed": False,
                "live_order_allowed": False,
                "output": str(output),
            }
        print(json.dumps(payload, indent=2, ensure_ascii=False))
        return 0
    except Exception as exc:
        payload = {
            "program": PROGRAM,
            "generated_utc": iso_utc(),
            "decision": "MACRO_DATA_INVENTORY_FAIL_CLOSED",
            "pass": False,
            "paper_order_allowed": False,
            "demo_order_allowed": False,
            "live_order_allowed": False,
            "error": f"{type(exc).__name__}: {exc}",
        }
        print(json.dumps(payload, indent=2, ensure_ascii=False))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
